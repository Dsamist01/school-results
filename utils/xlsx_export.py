import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from utils.grading import compute_class_results


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_class_workbook(school_class, term, session, school_name=""):
    data = compute_class_results(school_class, term, session)
    class_subjects = data["class_subjects"]
    rows = data["students"]

    wb = Workbook()
    ws = wb.active
    ws.title = school_class.name[:31] or "Class"

    title = f"{school_name} - {school_class.name} - {term} {session}".strip(" -")
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                    end_column=2 + len(class_subjects) * 6)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # Header rows: row3 = subject names (merged across 6 cols), row4 = sub-columns
    SUBCOLS = ["TEST", "EXAM", "TOTAL", "AVERAGE", "POSITION", "GRADE"]
    r_subject, r_subcol = 3, 4
    ws.cell(row=r_subject, column=1, value="S/N").font = Font(bold=True)
    ws.cell(row=r_subject, column=2, value="NAME OF STUDENT").font = Font(bold=True)
    ws.merge_cells(start_row=r_subject, start_column=1, end_row=r_subcol, end_column=1)
    ws.merge_cells(start_row=r_subject, start_column=2, end_row=r_subcol, end_column=2)
    for c in (1, 2):
        cell = ws.cell(row=r_subject, column=c)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col = 3
    for cs in class_subjects:
        ws.merge_cells(start_row=r_subject, start_column=col, end_row=r_subject, end_column=col + 5)
        cell = ws.cell(row=r_subject, column=col, value=cs.subject.name.upper())
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        for i, label in enumerate(SUBCOLS):
            sc = ws.cell(row=r_subcol, column=col + i, value=label)
            sc.fill = SUB_FILL
            sc.font = Font(bold=True, size=9)
            sc.alignment = Alignment(horizontal="center")
        col += 6

    # Data rows
    start_data_row = r_subcol + 1
    for i, row in enumerate(rows, start=1):
        r = start_data_row + i - 1
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=row["student"].name)
        col = 3
        for cs in class_subjects:
            sub = row["subjects"][cs.id]
            ws.cell(row=r, column=col, value=sub["test"])
            ws.cell(row=r, column=col + 1, value=sub["exam"])
            ws.cell(row=r, column=col + 2, value=sub["total"])
            ws.cell(row=r, column=col + 3, value=sub["average"])
            ws.cell(row=r, column=col + 4, value=sub["position"])
            ws.cell(row=r, column=col + 5, value=sub["grade"])
            col += 6

    last_col = 2 + len(class_subjects) * 6
    last_row = start_data_row + len(rows) - 1
    for r in range(r_subject, max(last_row, start_data_row) + 1):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).border = BORDER

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    for c in range(3, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.freeze_panes = ws.cell(row=start_data_row, column=3)

def build_score_import_template(school_class, subject, students, existing_scores, term, session):
    """Builds a simple Name/Test/Exam xlsx pre-filled with this class's
    students (and any existing scores) for the teacher to fill in and
    re-upload."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"

    ws.cell(row=1, column=1, value=f"{school_class.name} - {subject.name} - {term} {session}").font = \
        Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    headers = ["Name", "Test (max 40)", "Exam (max 60)"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")

    for i, student in enumerate(students, start=4):
        sc_obj = existing_scores.get(student.id)
        ws.cell(row=i, column=1, value=student.name)
        ws.cell(row=i, column=2, value=sc_obj.test if sc_obj else None)
        ws.cell(row=i, column=3, value=sc_obj.exam if sc_obj else None)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

