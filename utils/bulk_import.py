"""Parsing + name-matching logic for bulk score import (xlsx/csv upload)."""
import io
import csv
import difflib
from openpyxl import load_workbook


def _normalize(name):
    return " ".join((name or "").strip().lower().split())


def _find_col(headers, *candidates):
    """headers: list of lowercased header strings. Returns index or None."""
    for cand in candidates:
        for i, h in enumerate(headers):
            if cand in h:
                return i
    return None


def parse_score_file(file_storage):
    """Parses an uploaded .xlsx or .csv file and returns a list of
    {'name': str, 'test': float, 'exam': float} dicts.
    Expects columns containing 'name', 'test', and 'exam' in the header
    row (case-insensitive, any order, extra columns ignored)."""
    filename = (file_storage.filename or "").lower()
    rows = []

    if filename.endswith(".csv"):
        content = file_storage.read().decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(content))
        raw_rows = list(reader)
    else:
        wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True)
        ws = wb.active
        raw_rows = [[c.value for c in row] for row in ws.iter_rows()]

    raw_rows = [r for r in raw_rows if any(c not in (None, "") for c in r)]
    if not raw_rows:
        return []

    header = [str(c).strip().lower() if c is not None else "" for c in raw_rows[0]]
    name_idx = _find_col(header, "name")
    test_idx = _find_col(header, "test")
    exam_idx = _find_col(header, "exam")

    if name_idx is None:
        # fall back: assume first column is name if no header matched
        name_idx, test_idx, exam_idx = 0, 1, 2
        data_rows = raw_rows
    else:
        data_rows = raw_rows[1:]

    for r in data_rows:
        if name_idx >= len(r):
            continue
        name = str(r[name_idx]).strip() if r[name_idx] is not None else ""
        if not name:
            continue

        def _num(idx):
            if idx is None or idx >= len(r) or r[idx] in (None, ""):
                return 0
            try:
                return float(r[idx])
            except (ValueError, TypeError):
                return 0

        rows.append({"name": name, "test": _num(test_idx), "exam": _num(exam_idx)})

    return rows


def match_students(students, parsed_rows):
    """Matches parsed (name, test, exam) rows to existing Student records.

    Returns (matched, unmatched_file_rows, students_without_row)
      matched: {student_id: {'test':, 'exam':, 'matched_name':, 'exact': bool}}
      unmatched_file_rows: [{'name':, 'test':, 'exam':}] - names in file with no good match
      students_without_row: [Student] - students in class with no row in the file
    """
    by_norm = {_normalize(s.name): s for s in students}
    norm_names = list(by_norm.keys())

    matched = {}
    unmatched_file_rows = []
    used_student_ids = set()

    for row in parsed_rows:
        norm = _normalize(row["name"])
        student = by_norm.get(norm)
        exact = True
        if not student:
            close = difflib.get_close_matches(norm, norm_names, n=1, cutoff=0.82)
            if close:
                student = by_norm[close[0]]
                exact = False
        if student:
            matched[student.id] = {"test": row["test"], "exam": row["exam"],
                                    "matched_name": student.name, "exact": exact,
                                    "file_name": row["name"]}
            used_student_ids.add(student.id)
        else:
            unmatched_file_rows.append(row)

    students_without_row = [s for s in students if s.id not in used_student_ids]
    return matched, unmatched_file_rows, students_without_row
